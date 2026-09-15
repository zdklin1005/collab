from __future__ import annotations

from pathlib import Path
from xml.etree.ElementTree import Element, SubElement, tostring
from xml.dom import minidom

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = Path(__file__).resolve().parent
DRAWIO = OUT / 'LocalQuest_User_Management_UML.drawio'
DOCX = OUT / 'LocalQuest_User_Management_Use_Case_Descriptions_and_Functional_Requirements.docx'

ACTOR = 'shape=umlActor;verticalLabelPosition=bottom;verticalAlign=top;html=1;outlineConnect=0;'
OVAL = 'ellipse;whiteSpace=wrap;html=1;align=center;verticalAlign=middle;'
BOUNDARY = 'rounded=0;whiteSpace=wrap;html=1;movable=0;resizable=0;rotatable=0;deletable=0;editable=0;locked=1;connectable=0;'
TITLE = 'text;html=1;whiteSpace=wrap;strokeColor=none;fillColor=none;align=center;verticalAlign=middle;fontStyle=1;fontSize=16;'
LABEL = 'text;html=1;whiteSpace=wrap;strokeColor=none;fillColor=none;align=left;verticalAlign=middle;fontStyle=1;fontSize=12;'
ASSOC = 'edgeStyle=orthogonalEdgeStyle;rounded=0;orthogonalLoop=1;jettySize=auto;html=1;endArrow=none;endFill=0;'
INCLUDE = 'edgeStyle=orthogonalEdgeStyle;rounded=0;orthogonalLoop=1;jettySize=auto;html=1;dashed=1;endArrow=open;endFill=0;'
EXTEND = INCLUDE
LANE = 'swimlane;html=1;startSize=28;whiteSpace=wrap;fillColor=#f7f9fc;'
ACTION = 'rounded=1;whiteSpace=wrap;html=1;align=center;verticalAlign=middle;'
DECISION = 'strokeWidth=2;html=1;shape=mxgraph.flowchart.decision;whiteSpace=wrap;align=center;verticalAlign=middle;'
START = 'ellipse;fillColor=#000000;strokeColor=#000000;'
END = 'ellipse;shape=endState;fillColor=#000000;strokeColor=#000000;'
FLOW = 'edgeStyle=orthogonalEdgeStyle;rounded=0;orthogonalLoop=1;jettySize=auto;html=1;endArrow=block;endFill=1;'


def add_cell(root, cid, value='', style='', x=0, y=0, w=0, h=0, parent='1', vertex=True):
    attrs = {'id': cid, 'parent': parent, 'value': value, 'style': style}
    if vertex:
        attrs['vertex'] = '1'
    cell = SubElement(root, 'mxCell', attrs)
    SubElement(cell, 'mxGeometry', {'x': str(x), 'y': str(y), 'width': str(w), 'height': str(h), 'as': 'geometry'})
    return cell


def add_edge(root, cid, source, target, style=FLOW, label=None, points=None):
    cell = SubElement(root, 'mxCell', {
        'id': cid, 'parent': '1', 'source': source, 'target': target,
        'style': style, 'edge': '1',
    })
    geo = SubElement(cell, 'mxGeometry', {'relative': '1', 'as': 'geometry'})
    if points:
        arr = SubElement(geo, 'Array', {'as': 'points'})
        for x, y in points:
            SubElement(arr, 'mxPoint', {'x': str(x), 'y': str(y)})
    if label:
        label_cell = SubElement(root, 'mxCell', {
            'id': cid + '_label', 'parent': cid,
            'style': 'edgeLabel;html=1;align=center;verticalAlign=middle;resizable=0;points=[];fontSize=11;',
            'value': label, 'vertex': '1', 'connectable': '0',
        })
        lg = SubElement(label_cell, 'mxGeometry', {'relative': '1', 'as': 'geometry', 'x': '0', 'y': '-1'})
        SubElement(lg, 'mxPoint', {'as': 'offset'})
    return cell


def diagram(name, width=1200, height=900):
    d = Element('diagram', {'name': name, 'id': name.lower().replace(' ', '_').replace('-', '_')})
    model = SubElement(d, 'mxGraphModel', {
        'dx': '1200', 'dy': '900', 'grid': '1', 'gridSize': '10', 'guides': '1',
        'tooltips': '1', 'connect': '1', 'arrows': '1', 'fold': '1', 'page': '1',
        'pageScale': '1', 'pageWidth': str(width), 'pageHeight': str(height), 'math': '0', 'shadow': '0',
    })
    root = SubElement(model, 'root')
    SubElement(root, 'mxCell', {'id': '0'})
    SubElement(root, 'mxCell', {'id': '1', 'parent': '0'})
    return d, root


def overview_page():
    d, r = diagram('UC-OVERVIEW', 1200, 900)
    add_cell(r, 'b', '', BOUNDARY, 250, 70, 730, 720)
    add_cell(r, 'title', 'LocalQuest — User Management Module Overview', TITLE, 475, 80, 300, 28)
    add_cell(r, 'tourist', 'Tourist', ACTOR, 120, 280, 40, 70)
    add_cell(r, 'merchant', 'Merchant', ACTOR, 120, 590, 40, 70)
    items = [
        ('reg', 'Register Account', 360, 145), ('login', 'Sign In', 360, 245),
        ('reset', 'Recover Password', 360, 345), ('tourist_profile', 'Manage Tourist Profile', 360, 445),
        ('security', 'Manage Security &amp; Preferences', 360, 555), ('account', 'Sign Out / Delete Account', 360, 665),
        ('merchant_profile', 'Manage Merchant Profile', 690, 245),
        ('business', 'Manage Business Workspace', 690, 360),
        ('offers', 'Manage Campaigns &amp; Vouchers', 690, 500),
    ]
    for cid, text, x, y in items:
        add_cell(r, cid, text, OVAL, x, y, 190, 70)
    for n, target in enumerate(['reg', 'login', 'reset', 'tourist_profile', 'security', 'account']):
        add_edge(r, f't{n}', 'tourist', target, ASSOC)
    for n, target in enumerate(['reg', 'login', 'reset', 'security', 'account', 'merchant_profile', 'business', 'offers']):
        add_edge(r, f'm{n}', 'merchant', target, ASSOC)
    return d


def detail_page():
    d, r = diagram('UC-UM-DETAILED', 1500, 980)
    add_cell(r, 'b', '', BOUNDARY, 280, 60, 950, 820)
    add_cell(r, 'title', 'LocalQuest — Detailed User Management Use Cases', TITLE, 590, 70, 330, 28)
    add_cell(r, 'tourist', 'Tourist', ACTOR, 120, 245, 40, 70)
    add_cell(r, 'merchant', 'Merchant', ACTOR, 120, 610, 40, 70)
    add_cell(r, 'firebase', 'Firebase Authentication &amp; Firestore', ACTOR, 1320, 305, 40, 70)
    add_cell(r, 'cloudinary', 'Cloudinary', ACTOR, 1320, 630, 40, 70)
    usecases = [
        ('register', 'Register Account', 385, 125), ('validate_register', 'Validate Registration Details', 670, 125),
        ('verify_email', 'Send Email Verification', 965, 125),
        ('login', 'Sign In', 385, 245), ('authenticate', 'Authenticate Credentials', 670, 245),
        ('recover', 'Recover Password', 965, 245),
        ('tourist_profile', 'Manage Tourist Profile', 385, 385), ('profile_photo', 'Update Profile Photo', 670, 385),
        ('security', 'Manage Account Security &amp; Preferences', 965, 385),
        ('merchant_profile', 'Manage Merchant Profile', 385, 535), ('business', 'Manage Business Workspace', 670, 535),
        ('pin', 'Pin Business Address', 965, 535),
        ('ssm', 'Verify SSM Certificate', 670, 670), ('offers', 'Manage Campaigns &amp; Vouchers', 965, 670),
        ('validate_offer', 'Validate Offer Details', 965, 790), ('signout', 'Sign Out / Delete Account', 385, 790),
    ]
    for cid, text, x, y in usecases:
        add_cell(r, cid, text, OVAL, x, y, 205, 62)
    for n, target in enumerate(['register', 'login', 'recover', 'tourist_profile', 'security', 'profile_photo', 'signout']):
        add_edge(r, f't{n}', 'tourist', target, ASSOC)
    for n, target in enumerate(['register', 'login', 'recover', 'security', 'profile_photo', 'merchant_profile', 'business', 'offers', 'signout']):
        add_edge(r, f'm{n}', 'merchant', target, ASSOC)
    for n, target in enumerate(['register', 'login', 'recover', 'tourist_profile', 'security', 'merchant_profile', 'business', 'offers', 'signout']):
        add_edge(r, f'f{n}', 'firebase', target, ASSOC)
    for n, target in enumerate(['profile_photo', 'business', 'offers']):
        add_edge(r, f'c{n}', 'cloudinary', target, ASSOC)
    for cid, src, dst, kind in [
        ('i1', 'register', 'validate_register', 'include'), ('i2', 'register', 'verify_email', 'include'),
        ('i3', 'login', 'authenticate', 'include'), ('e1', 'profile_photo', 'tourist_profile', 'extend'),
        ('e2', 'profile_photo', 'merchant_profile', 'extend'), ('i4', 'merchant_profile', 'business', 'include'),
        ('e3', 'pin', 'business', 'extend'), ('e4', 'ssm', 'business', 'extend'),
        ('i5', 'offers', 'validate_offer', 'include'),
    ]:
        add_edge(r, cid, src, dst, INCLUDE if kind == 'include' else EXTEND, f'&lt;&lt;{kind}&gt;&gt;')
    return d


def activity_page(name, lanes, nodes, edges, width=1500, height=1000):
    d, r = diagram(name, width, height)
    lane_w = (width - 120) // len(lanes)
    for i, lane in enumerate(lanes):
        add_cell(r, f'lane{i}', lane, LANE, 60 + i * lane_w, 70, lane_w, height - 140)
    add_cell(r, 'title', name.replace('-', ' '), TITLE, width // 2 - 200, 25, 400, 28)
    for cid, typ, text, x, y, w, h in nodes:
        style = {'start': START, 'end': END, 'action': ACTION, 'decision': DECISION, 'note': LABEL}[typ]
        add_cell(r, cid, text, style, x, y, w, h)
    for idx, (src, dst, label, points) in enumerate(edges):
        add_edge(r, f'e{idx}', src, dst, FLOW, label, points)
    return d


def activities():
    ad1_nodes = [
        ('start', 'start', '', 245, 115, 22, 22), ('choose', 'action', 'Choose account type and\nauthentication action', 170, 165, 170, 55),
        ('choice', 'decision', '', 205, 255, 85, 65), ('reg_data', 'action', 'Enter registration\ndetails', 150, 365, 210, 55),
        ('login_data', 'action', 'Enter email/username\nand password', 150, 505, 210, 55), ('forgot', 'action', 'Request password reset', 150, 650, 210, 55),
        ('validate', 'action', 'Validate required fields\nand password policy', 600, 365, 210, 55), ('create', 'action', 'Create account and\nrole profile', 600, 445, 210, 55),
        ('email', 'action', 'Send verification email\nand open role home', 600, 525, 210, 55), ('throttle', 'action', 'Check 10-minute\nlogin throttle', 600, 505, 210, 55),
        ('auth', 'action', 'Authenticate credentials\nand load role profile', 600, 585, 210, 55), ('failure', 'decision', '', 665, 670, 85, 65),
        ('lock', 'action', 'Record failed attempt;\nlock after 5 failures', 600, 760, 210, 55), ('route', 'action', 'Route to tourist or\nmerchant interface', 1030, 555, 210, 55),
        ('reset_send', 'action', 'Send password-reset\nemail', 1030, 650, 210, 55), ('show_error', 'action', 'Show validation or\nlogin error', 1030, 760, 210, 55),
        ('end', 'end', '', 1125, 850, 22, 22),
    ]
    ad1_edges = [
        ('start','choose',None,None), ('choose','choice',None,None), ('choice','reg_data','[Register]',None), ('choice','login_data','[Sign in]',None), ('choice','forgot','[Forgot password]',None),
        ('reg_data','validate',None,None), ('validate','create','[Valid]',None), ('validate','show_error','[Invalid]',None), ('create','email',None,None), ('email','route',None,None),
        ('login_data','throttle',None,None), ('throttle','auth','[Not locked]',None), ('throttle','show_error','[Locked]',None), ('auth','failure',None,None), ('failure','route','[Valid]',None), ('failure','lock','[Invalid]',None), ('lock','show_error',None,None),
        ('forgot','reset_send',None,None), ('reset_send','end',None,None), ('route','end',None,None), ('show_error','end',None,None),
    ]
    ad2_nodes = [
        ('start','start','',245,115,22,22), ('open','action','Open Profile',170,165,170,55), ('choice','decision','',205,260,85,65),
        ('details','action','Edit name, username,\nphone or birthday',150,365,210,55), ('prefs','action','Change preferences,\nsecurity or linked sign-in',150,515,210,55),
        ('photo','action','Choose and crop\nprofile photo',150,665,210,55), ('delete','action','Confirm account deletion\nand enter current password',150,805,210,55),
        ('validate','action','Validate input and\ncurrent session',610,365,210,55), ('save','action','Save profile or\npreference changes',610,445,210,55),
        ('reauth','action','Re-authenticate before\nsecurity/deletion action',610,585,210,55), ('imagecheck','action','Validate photo file\nand upload request',610,705,210,55),
        ('remove','action','Remove owned data\nand account record',610,805,210,55), ('cloud','action','Store image and return\nsecure image URL',1035,705,210,55),
        ('persist','action','Persist user/profile\nchanges in Firestore',1035,445,210,55), ('success','action','Refresh profile and\nshow confirmation',1035,535,210,55),
        ('return','action','Return to role selection',1035,805,210,55), ('end','end','',1125,895,22,22),
    ]
    ad2_edges = [
        ('start','open',None,None),('open','choice',None,None),('choice','details','[Edit details]',None),('choice','prefs','[Security/preferences]',None),('choice','photo','[Photo]',None),('choice','delete','[Delete account]',None),
        ('details','validate',None,None),('validate','save','[Valid]',None),('validate','success','[Invalid: show error]',None),('save','persist',None,None),('persist','success',None,None),
        ('prefs','reauth',None,None),('reauth','persist','[Authorised]',None),('reauth','success','[Rejected: show error]',None),
        ('photo','imagecheck',None,None),('imagecheck','cloud','[Valid]',None),('imagecheck','success','[Invalid: show error]',None),('cloud','persist',None,None),
        ('delete','reauth',None,None),('reauth','remove','[Deletion confirmed]',[(715,650),(715,830)]),('remove','return',None,None),('success','end',None,None),('return','end',None,None),
    ]
    ad3_nodes = [
        ('start','start','',245,115,22,22),('open','action','Open Merchant Profile\nor Campaign workspace',150,165,210,55),('choice','decision','',205,260,85,65),
        ('business_form','action','Create or edit business\nworkspace details',150,365,210,55),('pin','action','Search, use current location\nor pin business address',150,445,210,55),
        ('cert','action','Optional: scan/upload\nSSM certificate',150,535,210,55),('offer_form','action','Create/edit campaign\nor voucher details',150,700,210,55),('review','action','Review and submit',150,780,210,55),
        ('validate_b','action','Validate business fields,\ncategory and dietary status',610,365,220,55),('geocode','action','Reverse-geocode selected\npin into address fields',610,445,220,55),
        ('ocr','action','Analyse SSM text, number,\nexpiry and name match',610,535,220,55),('validate_o','action','Validate owner, active\nbusiness, dates and limits',610,740,220,55),
        ('upload_b','action','Optional: upload business\nphoto/certificate',1035,455,220,55),('save_b','action','Save workspace and\nverification status',1035,565,220,55),
        ('upload_o','action','Optional: upload\ncampaign poster',1035,700,220,55),('save_o','action','Save offer and resolve\nactive/scheduled/inactive',1035,790,220,55),('end','end','',1135,890,22,22),
    ]
    ad3_edges = [
        ('start','open',None,None),('open','choice',None,None),('choice','business_form','[Business workspace]',None),('choice','offer_form','[Campaign/voucher]',None),
        ('business_form','validate_b',None,None),('validate_b','pin','[Valid]',None),('validate_b','end','[Invalid: show error]',None),('pin','geocode',None,None),('geocode','cert',None,None),('cert','ocr',None,None),('ocr','upload_b','[Continue]',None),('upload_b','save_b',None,None),('save_b','end',None,None),
        ('offer_form','review',None,None),('review','validate_o',None,None),('validate_o','upload_o','[Valid]',None),('validate_o','end','[Invalid: show error]',None),('upload_o','save_o',None,None),('save_o','end',None,None),
    ]
    return [
        activity_page('AD-UM-1 Account Registration, Sign-In and Recovery', ['User', 'LocalQuest', 'Firebase Authentication / Firestore'], ad1_nodes, ad1_edges),
        activity_page('AD-UM-2 Tourist Account and Profile Management', ['Tourist', 'LocalQuest', 'Cloudinary / Firebase'], ad2_nodes, ad2_edges),
        activity_page('AD-UM-3 Merchant Business and Offer Management', ['Merchant', 'LocalQuest', 'Cloudinary / Firebase / Address Service'], ad3_nodes, ad3_edges),
    ]


def build_drawio():
    root = Element('mxfile', {'host': 'app.diagrams.net', 'pages': '5', 'agent': 'LocalQuest user-management artefact'})
    for page in [overview_page(), detail_page(), *activities()]:
        root.append(page)
    raw = minidom.parseString(tostring(root, encoding='utf-8')).toprettyxml(indent='  ', encoding='UTF-8')
    DRAWIO.write_bytes(raw)


def set_cell_shading(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), color)
    tc_pr.append(shd)


def set_cell_border(cell, color='1F4E79', size='8'):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tc_pr.append(borders)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag = qn(f'w:{edge}')
        item = borders.find(tag)
        if item is None:
            item = OxmlElement(f'w:{edge}')
            borders.append(item)
        item.set(qn('w:val'),'single'); item.set(qn('w:sz'),size); item.set(qn('w:color'),color)


def cell_text(cell, text, bold=False, color=None, size=9):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    r.bold = bold
    r.font.name = 'Arial'
    r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.bold = True
    r.font.name = 'Arial'
    r.font.size = Pt(15 if level == 1 else 12)
    r.font.color.rgb = RGBColor(31, 78, 121)
    return p


def para(doc, text, bold=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(text)
    r.bold = bold
    r.font.name = 'Arial'
    r.font.size = Pt(9.5)
    return p


def use_case_table(doc, uc):
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.allow_autofit = False
    table.columns[0].width = Cm(4.0)
    table.columns[1].width = Cm(12.5)
    def row(label, value, header=False):
        cells = table.add_row().cells
        cell_text(cells[0], label, bold=True, color='FFFFFF' if header else '1F1F1F')
        cell_text(cells[1], value, bold=header, color='FFFFFF' if header else '1F1F1F')
        if header:
            set_cell_shading(cells[0], '1F4E79'); set_cell_shading(cells[1], '1F4E79')
        else:
            set_cell_shading(cells[0], 'D9EAF7')
        for c in cells: set_cell_border(c)
        return cells
    row('USE CASE NAME', f"{uc['id']}  {uc['name']}", header=True)
    row('USE CASE DESCRIPTION', uc['description'])
    row('PRE-CONDITION', uc['pre'])
    row('POST-CONDITION', uc['post'])
    cells = table.add_row().cells
    cells[0].merge(cells[1]); cell_text(cells[0], 'BASIC FLOW', bold=True, color='FFFFFF'); set_cell_shading(cells[0], '1F4E79'); set_cell_border(cells[0])
    cells = table.add_row().cells
    cell_text(cells[0], 'ACTOR', bold=True, color='FFFFFF'); cell_text(cells[1], 'SYSTEM', bold=True, color='FFFFFF')
    for c in cells: set_cell_shading(c, '5B9BD5'); set_cell_border(c)
    for actor, system in uc['basic']:
        cells = table.add_row().cells
        cell_text(cells[0], actor); cell_text(cells[1], system)
        for c in cells: set_cell_border(c)
    cells = table.add_row().cells
    cells[0].merge(cells[1]); cell_text(cells[0], 'ALTERNATE FLOW', bold=True, color='FFFFFF'); set_cell_shading(cells[0], '1F4E79'); set_cell_border(cells[0])
    for item in uc['alternate']:
        cells = table.add_row().cells; cells[0].merge(cells[1]); cell_text(cells[0], item)
        set_cell_shading(cells[0], 'F4F8FB'); set_cell_border(cells[0])
    cells = table.add_row().cells
    cells[0].merge(cells[1]); cell_text(cells[0], 'MESSAGE SECTION', bold=True, color='FFFFFF'); set_cell_shading(cells[0], '1F4E79'); set_cell_border(cells[0])
    for item in uc['messages']:
        cells = table.add_row().cells; cells[0].merge(cells[1]); cell_text(cells[0], item); set_cell_border(cells[0])
    cells = table.add_row().cells
    cells[0].merge(cells[1]); cell_text(cells[0], 'CONSTRAINT SECTION', bold=True, color='FFFFFF'); set_cell_shading(cells[0], '1F4E79'); set_cell_border(cells[0])
    for item in uc['constraints']:
        cells = table.add_row().cells; cells[0].merge(cells[1]); cell_text(cells[0], item); set_cell_border(cells[0])
    return table


def cases():
    return [
        dict(id='UC-UM01', name='Register Account', description='Creates a LocalQuest Tourist or Merchant account using email and password, and creates the role-specific profile. Merchant registration also creates the first business workspace.', pre='The user is not signed in and has selected Tourist or Merchant.', post='A Firebase Authentication account and Firestore user profile exist. A merchant also has an initial business workspace; an email verification request is sent where delivery is available.', basic=[('1. Selects Create an account and selects a role.','2. Displays the relevant Tourist or Merchant registration form.'),('3. Enters required details and submits the form.','4. Validates the email, password, phone and role-specific fields.'),('', '5. Creates the authentication account and role-based profile data.'),('', '6. Creates the first merchant business workspace when the selected role is Merchant.'),('', '7. Sends an email-verification request and opens the role-specific interface.')], alternate=['A1: Invalid or incomplete details. A1.1 The system identifies the affected field and remains on the form.','A2: Email already registered. A2.1 The system displays a message and directs the user to sign in or recover the password.','A3: Profile write fails after authentication creation. A3.1 The system discards the incomplete account where possible and reports the failure.'], messages=['M1: “Complete the required fields to create your LocalQuest account.”','M2: “Check your inbox to verify your email address.”'], constraints=['C1: Password must be 6–128 characters and must not be an obvious or repeated password.','C2: Tourist username must meet the app format rules; merchant business category and address are required.']),
        dict(id='UC-UM02', name='Sign In', description='Authenticates a Tourist or Merchant using an email address or saved username mapping, then opens the interface matching the account role.', pre='The user has a registered LocalQuest account and selected the correct role.', post='A role-appropriate authenticated session is established, or the failed-attempt counter is updated.', basic=[('1. Chooses Sign in for Tourist or Merchant.','2. Displays the sign-in form.'),('3. Enters email/username and password, then submits.','4. Resolves the identifier, checks the local lockout timer and authenticates credentials.'),('', '5. Loads the Firestore profile and verifies that its role matches the selected role.'),('', '6. Resets the failed-attempt counter, records the login time and opens the matching interface.')], alternate=['A1: Invalid credentials. A1.1 The system records the failure and displays an error.','A2: Fifth consecutive failed attempt. A2.1 The system locks further attempts for 10 minutes on that device.','A3: Selected role differs from the stored role. A3.1 The system signs out and asks the user to select the correct role.','A4: Profile document is missing. A4.1 The system attempts to recover a minimal role profile, then continues only if it can load it.'], messages=['M1: “Welcome back to LocalQuest.”','M2: “Incorrect email, username, or password.”','M3: “Too many failed attempts. Try again in X minutes.”'], constraints=['C1: Five consecutive failures trigger a 10-minute local lockout.','C2: The selected Tourist/Merchant role must equal the Firestore profile role.']),
        dict(id='UC-UM03', name='Recover Password', description='Sends a Firebase password-reset email to the registered email address.', pre='The user is on a LocalQuest authentication screen and knows the registered email address.', post='Firebase receives a password-reset request; the user can set a new password through the email flow.', basic=[('1. Selects Forgot password.','2. Displays the password-recovery form.'),('3. Enters the registered email address and submits.','4. Validates the email format and submits a password-reset request to Firebase Authentication.'),('', '5. Displays a confirmation and returns the user to the sign-in screen.')], alternate=['A1: Invalid email format. A1.1 The system requests a valid email address.','A2: Firebase rejects or cannot deliver the request. A2.1 The system displays the returned safe error message and allows retry.'], messages=['M1: “Enter the email linked to your LocalQuest account.”','M2: “If the account exists, password-reset instructions have been sent.”'], constraints=['C1: LocalQuest does not reveal whether an email address is registered in the confirmation message.']),
        dict(id='UC-UM04', name='Manage Tourist Profile', description='Allows a signed-in Tourist to view and update personal profile details, birthday, profile photo and notification/location preferences.', pre='The active account is authenticated with the Tourist role.', post='Valid profile and preference changes are persisted to the user profile; an uploaded photo URL is stored with the profile.', basic=[('1. Opens Profile.','2. Displays the current profile, statistics and profile actions.'),('3. Chooses an edit action.','4. Displays the relevant form or selector.'),('5. Updates allowed profile fields, preferences or photo and submits.','6. Validates input. For a photo, validates file type/size, allows crop, uploads it and receives a secure URL.'),('', '7. Stores the changed profile data in Firestore and refreshes the profile view.')], alternate=['A1: Invalid profile details. A1.1 The system identifies the invalid field and preserves the entered values.','A2: Unsupported/oversized image or failed upload. A2.1 The system retains the existing photo and shows an upload error.','A3: Permission is denied for a preference-dependent device feature. A3.1 The system explains that the preference cannot take effect until permission is granted.'], messages=['M1: “Profile updated.”','M2: “Choose a JPG, PNG or WEBP profile photo under 5 MB.”'], constraints=['C1: Only the signed-in user may edit that user profile.','C2: Profile images are limited to supported image formats and the app photo-size limit.']),
        dict(id='UC-UM05', name='Manage Account Security and Preferences', description='Lets a signed-in user manage password, email-change verification, optional biometric unlock, Google account linking, notification preferences and account deletion.', pre='The user is signed in. Sensitive changes require a current password or another suitable authentication method.', post='The selected security/preference change is applied, or the account and owned module records are removed after a confirmed deletion.', basic=[('1. Opens Settings / Password & Security.','2. Displays security providers, preferences and account actions.'),('3. Selects a security or preference action.','4. Displays the required form, device prompt or confirmation dialog.'),('5. Supplies the required current credential or consent.','6. Re-authenticates when required and applies the requested Firebase or local-device change.'),('', '7. Updates Firestore preference metadata where applicable and confirms the outcome.')], alternate=['A1: Current password is incorrect or stale. A1.1 The system rejects the sensitive change and asks the user to authenticate again.','A2: Google account is the only sign-in method. A2.1 The system prevents unlinking until a password or another sign-in method exists.','A3: Device biometrics are unavailable. A3.1 The system keeps biometric sign-in disabled and explains the limitation.','A4: Account deletion is cancelled. A4.1 The system keeps all records unchanged.'], messages=['M1: “Authentication is required before this change.”','M2: “Google is your only sign-in method; set a password before unlinking it.”','M3: “Your account and owned LocalQuest records will be permanently removed.”'], constraints=['C1: Password changes/deletion require re-authentication.','C2: Account deletion removes the user profile, owned businesses, campaigns and visited-place records before deleting the authentication account.']),
        dict(id='UC-UM06', name='Manage Merchant Profile', description='Allows a signed-in Merchant to view and update the merchant account profile, profile photo and account-level settings.', pre='The active account is authenticated with the Merchant role.', post='Valid merchant profile changes are saved and displayed in the merchant interface.', basic=[('1. Opens Merchant Profile.','2. Displays merchant account information, workspace access and settings actions.'),('3. Selects Edit profile or profile photo.','4. Displays the form or photo editor.'),('5. Updates details or confirms an edited photo.','6. Validates the data and uploads the photo when supplied.'),('', '7. Saves the account changes and refreshes the profile.')], alternate=['A1: Invalid text or contact information. A1.1 The system highlights the field and does not save.','A2: Photo upload fails. A2.1 The system keeps the current profile photo and lets the merchant retry.'], messages=['M1: “Merchant profile updated.”','M2: “Your existing profile photo was kept because the upload did not complete.”'], constraints=['C1: A merchant may update only the profile associated with the active authenticated account.']),
        dict(id='UC-UM07', name='Manage Business Workspace', description='Allows a Merchant to create, select, edit or delete their own business workspaces, including category-specific data, address search/current location/map pin, optional business photo and optional SSM verification.', pre='The active account is authenticated with the Merchant role.', post='The selected business workspace is saved with its owner ID, address/location data and verification status. Only the active selected workspace is used for offer management.', basic=[('1. Opens Business registration / workspace management.','2. Lists the merchant-owned workspaces and displays create/edit actions.'),('3. Selects a workspace action and enters business information.','4. Validates business name, category, address, phone and category-specific dietary certification where required.'),('5. Searches an address, uses current location or pins the business entrance.','6. Reverse-geocodes the pin and fills the displayed address fields; the merchant can adjust them.'),('7. Optionally adds a business photo or scans/uploads an SSM certificate.','8. Validates image input, analyses certificate text for SSM cues, registration number, expiry and business-name match, then determines verified/pending/rejected status.'),('', '9. Saves the owned workspace, location, optional photo URL and verification status in Firestore.')], alternate=['A1: Invalid business field. A1.1 The system shows the field-level validation message and does not save.','A2: Food/beverage category without dietary certification. A2.1 The system requires the merchant to choose a certification before saving.','A3: Location permission/network/address lookup fails. A3.1 The system allows manual address entry or retry.','A4: Certificate cannot be verified. A4.1 The workspace remains unverified or pending; it is not blocked from normal use.'], messages=['M1: “Pin the exact business entrance on the map.”','M2: “SSM verification is optional; unverified businesses remain usable without a verified badge.”','M3: “Business workspace saved.”'], constraints=['C1: A merchant may create/update/delete only workspaces whose ownerId matches the active user.','C2: A photo must meet supported type/size rules; a certificate analysis is prototype OCR/rule-based validation, not a live SSM registry query.']),
        dict(id='UC-UM08', name='Verify SSM Certificate', description='Optionally analyses an uploaded/captured Malaysian SSM certificate and applies a verified, pending-review or rejected verification result to a merchant business workspace.', pre='The merchant is editing a workspace they own and has an image or scan of an SSM certificate.', post='The system applies the detected registration number and verification status when the merchant confirms the result; the workspace may remain unverified.', basic=[('1. Selects Scan SSM registration certificate.','2. Lets the merchant capture an image or choose one from device storage.'),('3. Selects or captures the certificate.','4. Extracts available text and analyses Malaysian SSM keywords, registration number, date and business-name similarity.'),('', '5. Displays the detected number, confidence and verified/pending/rejected explanation.'),('6. Confirms application of the result.','7. Adds the result to the workspace form for saving.')], alternate=['A1: Image is unreadable/not an SSM-like document. A1.1 The system marks the result rejected and allows another image.','A2: Expiry date is detected as expired. A2.1 The system marks the result rejected.','A3: Evidence is incomplete. A3.1 The system marks the result pending review rather than showing a verified badge.'], messages=['M1: “Scan or choose a Malaysian SSM registration certificate.”','M2: “SSM certificate analysis is a prototype check and is not a live SSM registry confirmation.”'], constraints=['C1: The verified badge is shown only when the analysis meets the configured prototype verification threshold and detects a valid registration number.']),
        dict(id='UC-UM09', name='Manage Campaigns and Vouchers', description='Allows a Merchant to create, view, edit, pause, delete and associate promotional ads and vouchers for the currently selected active business workspace.', pre='The active account is authenticated as a Merchant and the selected business workspace belongs to that merchant and is active.', post='A valid campaign/voucher is stored under the selected business. Its effective status is active, scheduled or inactive according to dates and the merchant selection.', basic=[('1. Opens Campaigns & Vouchers and selects the active business workspace.','2. Loads that business’s ads and vouchers only.'),('3. Chooses Create or Edit and selects Promotional ad or Voucher.','4. Displays structured fields, preview/review and optional poster upload.'),('5. Enters title, description, dates, terms and voucher-specific value/quantity/limit fields, then reviews.','6. Validates ownership, active business state, text lengths, date range, status and voucher limits.'),('', '7. Uploads an optional poster and saves the offer in Firestore with the selected business ID.'),('', '8. Resolves the visible status: future start becomes Scheduled; past end becomes Inactive; valid current window can be Active.')], alternate=['A1: No active owned business is selected. A1.1 The system blocks save and requests a valid workspace.','A2: Invalid offer details. A2.1 The system shows the affected field messages and retains the form.','A3: Poster upload fails. A3.1 The system keeps the form and allows retry without saving an invalid URL.','A4: Merchant pauses/deletes an offer. A4.1 The system updates status or removes the selected owned offer.'], messages=['M1: “Select an active business before creating a campaign.”','M2: “Check voucher value, minimum spend and quantity limits.”','M3: “Campaign saved for the selected business.”'], constraints=['C1: The business ID must belong to the authenticated merchant and be active.','C2: Voucher percentage is >0 and ≤100; fixed amount, minimum spend, quantity and per-customer limits follow the app’s numeric limits.','C3: An expired campaign cannot be reactivated; a future campaign is scheduled.']),
    ]


def build_docx():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(1.5); sec.bottom_margin = Cm(1.5); sec.left_margin = Cm(1.45); sec.right_margin = Cm(1.45)
    normal = doc.styles['Normal']; normal.font.name = 'Arial'; normal.font.size = Pt(9.5)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('LOCALQUEST'); r.bold = True; r.font.name = 'Arial'; r.font.size = Pt(22); r.font.color.rgb = RGBColor(31,78,121)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('USER MANAGEMENT MODULE\nUse Case Descriptions and Functional Requirements'); r.bold = True; r.font.name = 'Arial'; r.font.size = Pt(15)
    para(doc, 'Prepared from the current LocalQuest Flutter implementation. Scope is intentionally limited to the User Management Module: account access, profiles, settings/security, merchant workspaces, optional SSM verification, and merchant campaigns/vouchers. Interactive map discovery, rewards, missions, reviews, social, chat and AI features are excluded.', False)
    heading(doc, 'B1. Use Case Diagram Scope', 1)
    para(doc, 'The companion draw.io file contains: (1) a user-management-only overview use case diagram, (2) a detailed use case diagram, and (3) three UML activity diagrams. The overview does not represent the other team modules.')
    heading(doc, 'B2. Use Case Descriptions', 1)
    for index, uc in enumerate(cases()):
        if index:
            if index == 7:
                # A section break avoids a Word pagination defect observed for
                # this table after a long preceding table.
                usecase_section = doc.add_section(WD_SECTION.NEW_PAGE)
                usecase_section.top_margin = Cm(1.5)
                usecase_section.bottom_margin = Cm(1.5)
                usecase_section.left_margin = Cm(1.45)
                usecase_section.right_margin = Cm(1.45)
            else:
                page_anchor = doc.add_paragraph()
                page_anchor.paragraph_format.page_break_before = True
                page_anchor.paragraph_format.space_after = Pt(0)
                page_anchor.paragraph_format.space_before = Pt(0)
                page_anchor.add_run(' ').font.size = Pt(1)
                spacer = doc.add_paragraph()
                spacer.paragraph_format.space_after = Pt(2)
                spacer.add_run(' ').font.size = Pt(4)
        use_case_table(doc, uc)
    doc.add_page_break()
    heading(doc, 'B3. Functional Requirements', 1)
    para(doc, 'The following functional requirements are extracted from the updated use cases. They describe implemented app behaviour; third-party delivery (email, Cloudinary upload and Firebase availability) still depends on the configured service being available at runtime.')
    requirements = [
        ('UC-UM01','FR-UM01','The system shall let a user choose Tourist or Merchant before registering.'),
        ('UC-UM01','FR-UM02','The system shall create a Firebase Authentication account from a valid email and password.'),
        ('UC-UM01','FR-UM03','The system shall create a role-specific Firestore user profile after registration.'),
        ('UC-UM01','FR-UM04','The system shall create an initial merchant business workspace during merchant registration.'),
        ('UC-UM01','FR-UM05','The system shall request email verification after successful registration.'),
        ('UC-UM02','FR-UM06','The system shall accept an email or resolvable username identifier for sign-in.'),
        ('UC-UM02','FR-UM07','The system shall route an authenticated user only to the interface matching the stored Tourist/Merchant role.'),
        ('UC-UM02','FR-UM08','The system shall record failed sign-in attempts and locally lock further attempts for 10 minutes after five consecutive failures.'),
        ('UC-UM03','FR-UM09','The system shall submit a password-reset email request for a valid email address.'),
        ('UC-UM04','FR-UM10','The system shall let a Tourist update allowed profile information and preferences.'),
        ('UC-UM04','FR-UM11','The system shall validate, crop and upload a supported profile photo, then store its returned URL with the profile.'),
        ('UC-UM05','FR-UM12','The system shall require re-authentication before changing a password, changing email or deleting an account.'),
        ('UC-UM05','FR-UM13','The system shall allow a user to enable/disable biometric app unlock when supported by the device.'),
        ('UC-UM05','FR-UM14','The system shall support linking/unlinking a Google sign-in provider without leaving the account with no usable sign-in method.'),
        ('UC-UM05','FR-UM15','The system shall delete the user profile and owned module records before deleting the Firebase Authentication account after confirmed deletion.'),
        ('UC-UM06','FR-UM16','The system shall let a Merchant update their account profile and profile photo.'),
        ('UC-UM07','FR-UM17','The system shall let a Merchant create, select, update and delete only their own business workspaces.'),
        ('UC-UM07','FR-UM18','The system shall support address suggestions, current-location retrieval and map pinning, then display a reverse-geocoded address rather than coordinates.'),
        ('UC-UM07','FR-UM19','The system shall require a halal/dietary certification selection for food and beverage business categories.'),
        ('UC-UM07','FR-UM20','The system shall store a business verification status of verified, pending review or rejected, while allowing unverified workspaces to remain usable.'),
        ('UC-UM08','FR-UM21','The system shall analyse a selected SSM certificate for expected text, registration number, expiry and business-name match to produce a prototype verification result.'),
        ('UC-UM09','FR-UM22','The system shall isolate campaigns and vouchers by their selected business workspace ID.'),
        ('UC-UM09','FR-UM23','The system shall validate campaign/voucher text, dates, voucher type, monetary values, quantity and per-customer limits before saving.'),
        ('UC-UM09','FR-UM24','The system shall permit an optional campaign poster upload and save the returned image URL only after a successful upload.'),
        ('UC-UM09','FR-UM25','The system shall resolve offer status as Active, Scheduled or Inactive from the chosen status and the start/end date window.'),
        ('UC-UM09','FR-UM26','The system shall prevent a merchant from saving an offer for a business that is not active or is not owned by the authenticated merchant.'),
    ]
    table = doc.add_table(rows=1, cols=3); table.alignment = WD_TABLE_ALIGNMENT.CENTER; table.autofit = False; table.allow_autofit = False
    widths = [Cm(3.0), Cm(3.0), Cm(10.5)]
    for c,w in zip(table.rows[0].cells,widths): c.width=w
    for c, text in zip(table.rows[0].cells, ['USE CASE','REQ. ID','REQ. STATEMENT']):
        cell_text(c,text,True,'FFFFFF');set_cell_shading(c,'1F4E79');set_cell_border(c)
    for uc, rid, statement in requirements:
        cells=table.add_row().cells
        for c, width, text in zip(cells, widths, [uc,rid,statement]):
            c.width = width
            cell_text(c,text)
            set_cell_border(c)
    section = doc.sections[0]
    footer = section.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer.add_run('LocalQuest — User Management Module'); fr.font.name='Arial'; fr.font.size=Pt(8); fr.font.color.rgb=RGBColor(100,100,100)
    doc.save(DOCX)


if __name__ == '__main__':
    OUT.mkdir(parents=True, exist_ok=True)
    build_drawio()
    build_docx()
    print(DRAWIO)
    print(DOCX)
